"""Заполнение рабочей тетради пары 2 (Ермилов Д.Е., вариант 27).
Использует результаты запросов из results.json (их собирает run_queries.py)."""
import json
import sys
from copy import deepcopy
from pathlib import Path

import docx
from docx.shared import RGBColor

HERE = Path(__file__).parent
SRC = HERE / "rt2.docx"
DST = HERE / "РТ (УП11) пара 2 ст Ермилов Д.Е..docx"
RES = json.loads((HERE / "results.json").read_text(encoding="utf-8"))
FIO = "Ермилов Даниил Евгеньевич"
BLUE = RGBColor(0x1F, 0x4E, 0xA8)

d = docx.Document(str(SRC))
P = d.paragraphs
T = d.tables


def check(par_or_cell):
    """Заменить все ☐ на ☑ в абзаце или ячейке."""
    pars = par_or_cell.paragraphs if hasattr(par_or_cell, "paragraphs") else [par_or_cell]
    for p in pars:
        for r in p.runs:
            r.text = r.text.replace("☐", "☑")


def fill_blank(p, value):
    """Заменить '_____' в абзаце на значение."""
    for r in p.runs:
        if "___" in r.text:
            r.text = r.text.split("_")[0] + value
            r.font.color.rgb = BLUE
            return
    add(p, " " + value)


def add(p, text):
    r = p.add_run(text)
    r.font.color.rgb = BLUE
    return r


def set_cell(cell, text):
    p = cell.paragraphs[0]
    for r in p.runs[1:]:
        r._element.getparent().remove(r._element)
    if p.runs:
        p.runs[0].text = text
        p.runs[0].font.color.rgb = BLUE
    else:
        add(p, text)


def set_code(p, code):
    """Абзац 'sql' -> текст кода с переносами строк."""
    p.runs[0].text = ""
    lines = code.strip("\n").split("\n")
    for i, line in enumerate(lines):
        r = add(p, line)
        r.font.name = "Consolas"
        if i < len(lines) - 1:
            r.add_break()


def answer_after(p, text):
    """Вставить новый абзац с ответом сразу после p."""
    new = deepcopy(p._element)
    p._element.addnext(new)
    np = docx.text.paragraph.Paragraph(new, p._parent)
    for r in np.runs[1:]:
        r._element.getparent().remove(r._element)
    np.runs[0].text = "Ответ: " + text
    np.runs[0].bold = False
    np.runs[0].font.color.rgb = BLUE
    return np


def fill_table(tbl, rows):
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            set_cell(tbl.rows[i].cells[j], str(val))


# --- Шапка ---
fill_blank(P[2], FIO)

# --- Задание 1 ---
set_cell(T[5].rows[1].cells[1], "☑")
fill_blank(P[22], "PostgreSQL 17")
fill_blank(P[23], "объектно-реляционная СУБД с открытым исходным кодом, строгая поддержка стандарта SQL, "
                  "CHECK/FOREIGN KEY, удобный pgAdmin 4 в комплекте; подходит для БД «Парфюмерия».")

# --- Задание 2 (вариант А: PostgreSQL) ---
for ti in (6, 7, 11):
    for row in T[ti].rows[1:]:
        check(row.cells[1])
check(P[33])
P[42].runs[0].text = "Результат (вариант Б не выполнялся, выбран PostgreSQL): "

# --- Задание 3 (вариант А: pgAdmin) ---
for ti in (18, 21):
    for row in T[ti].rows[1:]:
        check(row.cells[1])
for row in T[19].rows[1:3]:
    check(row.cells[1])
check(P[49])
P[55].runs[0].text = "Результат (вариант Б не выполнялся, выбран pgAdmin): "

# --- Задание 4 ---
fill_blank(P[61], RES["version"])
fill_blank(P[65], RES["date"])
fill_blank(P[69], RES["time"])
fill_blank(P[73], RES["result"])
fill_blank(P[77], RES["greeting"])
for i in (83, 99, 103, 116, 123, 139, 143, 147, 152):
    check(P[i])
fill_blank(P[91], RES["databases"])
fill_table(T[26], RES["q12"])
fill_table(T[27], RES["q13"])
fill_table(T[28], RES["q14"])

# --- Задание 5 ---
LEVEL1_TABLE = """CREATE TABLE products (
    id SERIAL PRIMARY KEY,
    name VARCHAR(100) NOT NULL,
    price NUMERIC(10,2) CHECK (price > 0)
);"""
LEVEL1_INSERT = """INSERT INTO products (name, price) VALUES
('Chanel No 5, 50 мл', 14500.00),
('Dior Sauvage, 100 мл', 12900.00),
('Lancome La Vie Est Belle, 75 мл', 11200.00);"""
LEVEL2_DB = """CREATE DATABASE library_db;
\\c library_db"""
LEVEL2_TABLE = """CREATE TABLE books (
    id SERIAL PRIMARY KEY,
    title VARCHAR(200) NOT NULL,
    author VARCHAR(100) NOT NULL,
    year INTEGER CHECK (year > 1900),
    price NUMERIC(10,2) CHECK (price > 0)
);"""
LEVEL2_INSERT = """INSERT INTO books (title, author, year, price) VALUES
('Парфюмер. История одного убийцы', 'Патрик Зюскинд', 1985, 450.00),
('Мастер и Маргарита', 'Михаил Булгаков', 1967, 520.00),
('Гарри Поттер и философский камень', 'Джоан Роулинг', 1997, 780.00),
('Код да Винчи', 'Дэн Браун', 2003, 610.00),
('Марсианин', 'Энди Вейер', 2011, 490.00);"""
LEVEL3_DB = """CREATE DATABASE company_db;
\\c company_db"""
LEVEL3_TABLES = """CREATE TABLE departments (
    id SERIAL PRIMARY KEY,
    name VARCHAR(100) NOT NULL
);
CREATE TABLE employees (
    id SERIAL PRIMARY KEY,
    name VARCHAR(100) NOT NULL,
    salary NUMERIC(10,2) CHECK (salary > 0),
    department_id INTEGER REFERENCES departments(id)
);"""
LEVEL3_INSERT = """INSERT INTO departments (name) VALUES ('IT'), ('HR');
INSERT INTO employees (name, salary, department_id) VALUES
('Иванов', 80000, (SELECT id FROM departments WHERE name = 'IT')),
('Петров', 90000, (SELECT id FROM departments WHERE name = 'IT')),
('Сидорова', 70000, (SELECT id FROM departments WHERE name = 'HR'));"""

set_code(P[161], LEVEL1_TABLE)
set_code(P[165], LEVEL1_INSERT)
set_code(P[170], LEVEL2_DB)
set_code(P[174], LEVEL2_TABLE)
set_code(P[178], LEVEL2_INSERT)
set_code(P[185], LEVEL3_DB)
set_code(P[191], LEVEL3_TABLES)
set_code(P[195], LEVEL3_INSERT)

for r, sql in zip(T[31].rows[2:], [
        "SELECT * FROM books WHERE year > 2000;",
        "SELECT * FROM books ORDER BY price;",
        "SELECT * FROM books WHERE price > 500;"]):
    set_cell(r.cells[2], sql)
for r, sql in zip(T[35].rows[1:], [
        "SELECT * FROM employees;",
        "SELECT * FROM employees WHERE salary > 80000;",
        "SELECT * FROM employees ORDER BY salary DESC;",
        "SELECT e.* FROM employees e JOIN departments d ON d.id = e.department_id WHERE d.name = 'IT';",
        "SELECT ROUND(AVG(salary), 2) AS avg_salary FROM employees;"]):
    set_cell(r.cells[2], sql)
for r, sql in zip(T[36].rows[1:], [
        "UPDATE employees SET salary = salary + 5000 WHERE name = 'Иванов';",
        "DELETE FROM employees WHERE name = 'Сидорова';"]):
    set_cell(r.cells[2], sql)

for i in (158, 162, 166, 171, 175, 179, 186, 192, 196):
    check(P[i])
for i, key in ((181, "l2"), (198, "l3q"), (200, "l3u")):
    check(P[i])
    add(P[i], " " + RES[key])
for i, key in ((158, "l1db"), (166, "l1"),):
    add(P[i], " " + RES[key])

# --- Контрольные вопросы (вставляем с конца, чтобы не сбить индексы) ---
ANSWERS = [
    "программное обеспечение для создания, хранения, обработки и управления базами данных.",
    "PostgreSQL — объектно-реляционная СУБД (порт 5432, pgAdmin, строгое следование стандарту SQL, "
    "сложные типы данных); MySQL — реляционная СУБД (порт 3306, MySQL Workbench), проще, популярна в веб-приложениях.",
    "графический инструмент для администрирования PostgreSQL: подключение к серверу, просмотр объектов БД, "
    "выполнение SQL-запросов (Query Tool).",
    "5432.",
    "3306.",
    "SELECT version();",
    "CREATE DATABASE имя_бд;",
    "SELECT * FROM имя_таблицы;",
    "SELECT ... WHERE условие;",
    "SELECT ... ORDER BY поле [ASC|DESC];",
    "INSERT INTO таблица (поля) VALUES (...);",
    "UPDATE таблица SET поле = значение WHERE условие;",
    "DELETE FROM таблица WHERE условие;",
    "первичный ключ — поле (набор полей), однозначно идентифицирующее запись: уникально и NOT NULL.",
    "внешний ключ — поле, ссылающееся на первичный ключ другой таблицы; обеспечивает ссылочную целостность "
    "(пример: employees.department_id → departments.id).",
]
REFLECTION = {
    224: "синтаксис psql-команд (\\c, \\l) и связь таблиц через FOREIGN KEY.",
    223: "задание повышенного уровня — выборка сотрудников отдела IT через JOIN по внешнему ключу.",
    222: "установка PostgreSQL и простые запросы SELECT, INSERT, UPDATE, DELETE.",
}
for i, text in REFLECTION.items():
    answer_after(P[i], text)
for i in range(217, 202, -1):
    answer_after(P[i], ANSWERS[i - 203])

# --- Итоги ---
for ti in (37, 38):
    for row in T[ti].rows[1:]:
        check(row.cells[2])
check(P[226])

# --- Самооценка ---
for row in T[39].rows[1:16]:
    set_cell(row.cells[3], "2")
set_cell(T[39].rows[16].cells[3], "30")

d.save(str(DST))
print("saved")
